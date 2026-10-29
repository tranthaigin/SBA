import os
import sys
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

from generate_docs_helper import (
    set_cell_background, set_cell_margins, add_header_banner,
    add_callout, add_detailed_step, add_summary_step,
    PRIMARY_COLOR, SECONDARY_COLOR, ACCENT_COLOR, DARK_TEXT, MUTED_TEXT, SUCCESS_COLOR
)

from steps_a1_data import a1_steps

# Define A2 Steps
a2_steps = [
    # Nhóm 1: Độc giả công khai
    (
        "noteTT/screenshots_a2/01_a2_public_news_list.png",
        "Chức năng 1.1: Giao diện độc giả đọc tin tức công khai (Public News Reading - /news)",
        "Độc giả truy cập /news mà không cần đăng nhập. Frontend gửi GET request tới /api/news/active. Backend lọc bài viết có newsStatus = true và trả về danh sách JSON.",
        "Trang hiển thị danh sách các bài viết tin tức đang hoạt động (Active), kèm thẻ Badge danh mục, tiêu đề, tóm tắt và nút 'Read', góc trên hiển thị nút 'Login (Staff / Admin)'.",
        "backend/.../NewsArticleController.java (Dòng 29-32) & frontend/src/pages/NewsList.jsx (Dòng 32-46)",
        "Phương thức getActive() trong NewsArticleController và fetchNews() trong NewsList.jsx: Cung cấp API và hiển thị tin tức công khai cho độc giả vãng lai.",
        "Độc giả chưa đăng nhập truy cập được danh sách tin tức Active, giao diện có nút Read và nút Login."
    ),
    (
        "noteTT/screenshots_a2/02_a2_public_news_search.png",
        "Chức năng 1.2: Độc giả tìm kiếm tin tức công khai theo từ khóa",
        "Nhập từ khóa 'student' vào ô tìm kiếm. Frontend lọc thời gian thực danh sách tin tức active theo tiêu đề và headline mà không cần gửi lại request server.",
        "Giao diện lọc mượt mà, chỉ giữ lại các bài viết có tiêu đề hoặc tóm tắt chứa từ khóa 'student'.",
        "frontend/src/pages/NewsList.jsx (Dòng 36-40)",
        "Bộ lọc client-side trên mảng res.data: Tìm kiếm tức thì tin tức công khai theo từ khóa nhập vào.",
        "Tìm kiếm tin tức công khai hoạt động chính xác, chỉ hiển thị bài báo chứa từ khóa phù hợp."
    ),

    # Nhóm 2: Đăng nhập & Xác thực
    (
        "noteTT/screenshots_a2/03_a2_login_modal.png",
        "Chức năng 2.1: Mở popup hộp thoại đăng nhập hệ thống (LoginModal)",
        "Bấm nút 'Login (Staff / Admin)', modal popup mở ra giữa màn hình với form email, password và 2 nút hỗ trợ điền nhanh thông tin tài khoản demo cho Staff và Admin.",
        "Xuất hiện hộp thoại 'Login to FUNews' với các nút 'Fill Staff Credentials', 'Fill Admin Credentials', ô Email Address và Password rõ ràng.",
        "frontend/src/components/LoginModal.jsx (Dòng 5-82)",
        "Component LoginModal và state showLogin trong App.jsx: Quản lý hộp thoại đăng nhập xác thực tài khoản.",
        "Popup đăng nhập hiển thị đầy đủ form nhập liệu và các phím tắt điền tài khoản mẫu tiện lợi."
    ),
    (
        "noteTT/screenshots_a2/04_a2_login_validation_error.png",
        "Chức năng 2.2: Bắt lỗi khi nhập sai mật khẩu đăng nhập",
        "Nhập email staff@funews.com và mật khẩu sai 'wrongpass123' rồi submit. Backend kiểm tra không khớp mật khẩu, ném BusinessRuleViolationException trả về HTTP 400.",
        "Xuất hiện thông báo cảnh báo lỗi màu đỏ 'Invalid email or password' ngay trong modal đăng nhập.",
        "backend/.../service/AuthService.java (Dòng 24-33) & frontend/src/components/LoginModal.jsx (Dòng 12-24)",
        "Phương thức login trong AuthService: Kiểm tra email và mật khẩu trong cơ sở dữ liệu, chặn đăng nhập sai.",
        "Hệ thống từ chối đăng nhập với mật khẩu sai và hiển thị thông báo lỗi trực quan."
    ),
    (
        "noteTT/screenshots_a2/05_a2_staff_login_success.png",
        "Chức năng 2.3: Đăng nhập thành công với vai trò Nhân viên (Staff - Role 2)",
        "Nhập đúng email staff@funews.com và mật khẩu @1. Backend xác thực thành công, lưu accountId vào HttpSession và trả về LoginResponse. Frontend cập nhật AuthContext.",
        "Modal đóng lại, Navbar hiển thị tên 'Staff Account', badge màu xanh 'Staff', và hiển thị thêm 3 menu quản lý: Categories, My Profile, My News History.",
        "backend/.../controller/AuthController.java (Dòng 34-39) & frontend/src/components/AppNavBar.jsx (Dòng 26-36)",
        "Phương thức login trong AuthController: Cấp phát phiên làm việc HttpSession cho nhân viên Staff.",
        "Đăng nhập thành công vai trò Staff, Navbar mở rộng thêm các menu quản lý chuyên quyền."
    ),

    # Nhóm 3: Quản lý Thể loại (Staff)
    (
        "noteTT/screenshots_a2/06_a2_category_list.png",
        "Chức năng 3.1: Hiển thị danh sách Thể loại trong cơ sở dữ liệu SQL Server (Category List)",
        "Staff bấm vào menu 'Categories' (/categories). Frontend gửi GET /api/categories. Interceptor kiểm tra session và role=2, cho phép gọi CategoryService.getAllCategories().",
        "Bảng hiển thị toàn bộ danh sách thể loại từ SQL Server gồm các cột: ID, Name, Description, Parent, Active (Yes/No), và 2 nút Edit, Delete.",
        "backend/.../controller/CategoryController.java (Dòng 24-27) & frontend/src/pages/CategoryList.jsx (Dòng 13-24)",
        "Phương thức getAll trong CategoryController và hàm load trong CategoryList.jsx: Nạp danh sách thể loại từ database.",
        "Bảng thể loại hiển thị đầy đủ thông tin từ SQL Server kèm các nút thao tác Edit và Delete."
    ),
    (
        "noteTT/screenshots_a2/07_a2_category_create_modal.png",
        "Chức năng 3.2: Mở Popup Modal thêm mới Thể loại (Create Category Modal)",
        "Staff bấm nút 'Create Category'. Modal popup CategoryForm mở ra với các trường: Category Name, Description, Parent Category ID (tùy chọn) và switch Active.",
        "Modal hiển thị tiêu đề 'Create Category', các ô nhập liệu sẵn sàng cho người dùng điền thông tin.",
        "frontend/src/components/CategoryForm.jsx (Dòng 80-148)",
        "Component CategoryForm: Xây dựng form nhập liệu thêm mới và chỉnh sửa thể loại.",
        "Modal thêm mới thể loại xuất hiện trên nền mờ với form nhập liệu chuẩn mực."
    ),
    (
        "noteTT/screenshots_a2/08_a2_category_create_success.png",
        "Chức năng 3.3: Thêm mới Thể loại thành công vào SQL Server",
        "Nhập Name 'Artificial Intelligence', Description và bấm 'Save'. Request POST gửi tới /api/categories. Backend JPA lưu thực thể Category vào database và sinh ID tự tăng.",
        "Modal đóng lại, dòng thể loại mới xuất hiện ngay trong bảng với tên 'Artificial Intelligence' và Active 'Yes'.",
        "backend/.../service/CategoryService.java (Dòng 42-49) & backend/.../controller/CategoryController.java (Dòng 39-42)",
        "Phương thức createCategory trong CategoryService: Lưu bản ghi thể loại mới vào SQL Server bằng JPA Repository.",
        "Thể loại mới được lưu thành công vào SQL Server và hiển thị ngay trên bảng danh sách."
    ),
    (
        "noteTT/screenshots_a2/09_a2_category_update_modal.png",
        "Chức năng 3.4: Mở Popup Modal chỉnh sửa Thể loại",
        "Bấm nút 'Edit' tại dòng thể loại vừa tạo. Modal mở ra với tiêu đề 'Edit Category' và điền sẵn thông tin cũ. Người dùng đổi tên thành 'Artificial Intelligence & Data'.",
        "Modal hiển thị đúng thông tin của thể loại được chọn, nút lưu hiển thị 'Save'.",
        "frontend/src/components/CategoryForm.jsx (Dòng 16-30)",
        "useEffect trong CategoryForm: Nạp dữ liệu của thể loại cần sửa vào form binding.",
        "Modal chỉnh sửa thể loại nạp chính xác dữ liệu cũ cho phép cập nhật."
    ),
    (
        "noteTT/screenshots_a2/10_a2_category_update_success.png",
        "Chức năng 3.5: Cập nhật Thể loại thành công",
        "Bấm 'Save', request PUT gửi tới /api/categories/{id}. Backend cập nhật các trường và lưu vào database.",
        "Bảng hiển thị tên mới 'Artificial Intelligence & Data', thông tin cập nhật tức thì trong SQL Server.",
        "backend/.../service/CategoryService.java (Dòng 51-58)",
        "Phương thức updateCategory trong CategoryService: Cập nhật thông tin thể loại theo ID.",
        "Thể loại cập nhật tên mới thành công, dữ liệu đồng bộ ngay lập tức."
    ),
    (
        "noteTT/screenshots_a2/11_a2_category_delete_prevent_in_use.png",
        "Chức năng 3.6: Chặn xóa Thể loại đang có bài viết liên kết (Ràng buộc nghiệp vụ)",
        "Bấm 'Delete' tại thể loại ID 1 (đang có bài viết tin tức liên kết). Backend gọi newsArticleRepository.existsByCategory(cat). Vì tồn tại bài viết, service ném BusinessRuleViolationException.",
        "Phía trên bảng xuất hiện Alert màu đỏ với thông báo: 'Cannot delete Category because it is associated with existing news articles'. Thể loại không bị xóa.",
        "backend/.../service/CategoryService.java (Dòng 60-66)",
        "Phương thức deleteCategory trong CategoryService: Kiểm tra ràng buộc toàn vẹn dữ liệu trước khi xóa.",
        "Hệ thống kích hoạt quy tắc chặn xóa thể loại đang có bài viết, bảo toàn tính toàn vẹn dữ liệu."
    ),
    (
        "noteTT/screenshots_a2/12_a2_category_delete_success.png",
        "Chức năng 3.7: Xóa Thể loại thành công khỏi SQL Server",
        "Bấm 'Delete' tại thể loại 'Artificial Intelligence & Data' (không có bài viết nào liên kết). Xác nhận hộp thoại, backend xóa bản ghi khỏi database.",
        "Dòng thể loại biến mất hoàn toàn khỏi bảng, cơ sở dữ liệu SQL Server được cập nhật sạch sẽ.",
        "backend/.../service/CategoryService.java (Dòng 65) & frontend/src/pages/CategoryList.jsx (Dòng 31-39)",
        "Phương thức remove trong CategoryList.jsx: Gọi API xóa thể loại hợp lệ khỏi database.",
        "Thể loại chưa có bài viết được xóa thành công sau khi xác nhận hộp thoại."
    ),
    (
        "noteTT/screenshots_a2/13_a2_category_search.png",
        "Chức năng 3.8: Tìm kiếm Thể loại theo từ khóa",
        "Nhập 'tech' vào ô tìm kiếm. Frontend gửi query ?keyword=tech, backend lọc bằng repository method findByCategoryNameContainingIgnoreCase.",
        "Bảng chỉ hiển thị các thể loại có tên chứa từ khóa 'tech' (ví dụ: Technology).",
        "backend/.../repository/CategoryRepository.java (Dòng 11) & backend/.../service/CategoryService.java (Dòng 26-31)",
        "Phương thức findByCategoryNameContainingIgnoreCase trong Spring Data JPA: Tìm kiếm không phân biệt hoa thường.",
        "Tìm kiếm thể loại theo từ khóa hoạt động chính xác qua câu truy vấn SQL Server."
    ),

    # Nhóm 4: Quản lý Tin tức (Staff)
    (
        "noteTT/screenshots_a2/14_a2_news_management_list.png",
        "Chức năng 4.1: Giao diện Quản lý Tin tức của Staff (News Management)",
        "Staff truy cập /news, hiển thị giao diện biên tập với nút '+ Create News Article', danh sách bài viết hiển thị cả tin Active lẫn Inactive kèm nút Edit, Delete cho từng bài.",
        "Xuất hiện nút '+ Create News Article' phía trên, mỗi thẻ tin tức có thêm 2 nút thao tác 'Edit' và 'Delete'.",
        "frontend/src/pages/NewsList.jsx (Dòng 117-123 & 175-184)",
        "Điều kiện render isStaff = user?.accountRole === 2: Mở khóa các công cụ biên tập tin tức cho nhân viên.",
        "Staff nhìn thấy đầy đủ nút tạo bài viết mới và các công cụ sửa/xóa trên từng tin tức."
    ),
    (
        "noteTT/screenshots_a2/15_a2_news_create_modal_tags.png",
        "Chức năng 4.2: Mở Popup Modal thêm mới Tin tức kèm chọn Thẻ Tags",
        "Bấm '+ Create News Article'. Modal mở ra cho phép nhập News Title, Headline, Category, chọn các Tags (#Spring, #React) dạng nút bấm, News Content, News Source và Active switch.",
        "Modal hiển thị đầy đủ các trường nhập liệu và danh sách nút bấm Tags trực quan, các tag được chọn chuyển sang màu xanh nổi bật.",
        "frontend/src/pages/NewsList.jsx (Dòng 230-310)",
        "Hàm handleTagToggle trong NewsList.jsx: Quản lý danh sách tagIds được chọn trên form.",
        "Modal tạo tin tức cho phép nhập đầy đủ nội dung và chọn nhiều Thẻ Tags trực quan."
    ),
    (
        "noteTT/screenshots_a2/16_a2_news_create_success.png",
        "Chức năng 4.3: Thêm mới bài viết Tin tức thành công vào cơ sở dữ liệu",
        "Bấm 'Create Article', request POST gửi tới /api/news. Backend tự động lấy accountId của Staff từ session để gán vào createdBy, gán ngày tạo LocalDateTime.now(), liên kết bảng quan hệ nhiều-nhiều với Tags.",
        "Bài viết mới 'Cloud Computing Workshop 2026' xuất hiện ngay trong danh sách với đầy đủ Category, badge trạng thái, các thẻ Tags và tên tác giả Staff.",
        "backend/.../service/NewsArticleService.java (Dòng 58-86)",
        "Phương thức createNews trong NewsArticleService: Ánh xạ quan hệ Many-to-Many Tags và One-to-Many Author tự động.",
        "Bài viết mới kèm danh sách thẻ Tags và tác giả được lưu thành công vào SQL Server."
    ),
    (
        "noteTT/screenshots_a2/17_a2_news_detail_read_modal.png",
        "Chức năng 4.4: Hộp thoại xem chi tiết nội dung bài viết tin tức (Read Detail Modal)",
        "Bấm nút 'Read' trên bài viết. Modal hiển thị đầy đủ toàn văn nội dung bài báo, headline, danh mục, nguồn tin và thông tin chi tiết tác giả (Tên & Email).",
        "Modal chi tiết mở lên trang trọng, nội dung giữ nguyên định dạng ngắt dòng và hiển thị đầy đủ thông tin nguồn tin/tác giả.",
        "frontend/src/pages/NewsList.jsx (Dòng 198-228)",
        "State viewArticle và Modal chi tiết bài báo: Cho phép người dùng đọc toàn văn nội dung bài viết.",
        "Hộp thoại đọc bài báo hiển thị đầy đủ nội dung chi tiết, nguồn tin và tác giả xuất bản."
    ),
    (
        "noteTT/screenshots_a2/18_a2_news_update_modal.png",
        "Chức năng 4.5: Mở Popup Modal chỉnh sửa bài viết Tin tức",
        "Bấm nút 'Edit' tại bài viết vừa tạo. Modal nạp toàn bộ thông tin cũ, cho phép người dùng chỉnh sửa tiêu đề thành 'Cloud & DevOps Summit 2026 - Extended' và thay đổi nội dung.",
        "Modal hiển thị tiêu đề 'Edit News Article' với các trường thông tin điền sẵn sẵn sàng cập nhật.",
        "frontend/src/pages/NewsList.jsx (Dòng 76-92)",
        "Hàm handleOpenEdit: Nạp dữ liệu bài viết vào state formData để chuẩn bị cập nhật.",
        "Modal chỉnh sửa bài viết nạp đầy đủ thông tin cũ, cho phép cập nhật tiêu đề và nội dung."
    ),
    (
        "noteTT/screenshots_a2/19_a2_news_update_success.png",
        "Chức năng 4.6: Cập nhật bài viết Tin tức thành công",
        "Bấm 'Save Changes', request PUT gửi lên backend. Backend cập nhật nội dung và tự động cập nhật modifiedDate = LocalDateTime.now().",
        "Thẻ bài viết trên giao diện đổi sang tiêu đề mới 'Cloud & DevOps Summit 2026 - Extended', dữ liệu được cập nhật thành công.",
        "backend/.../service/NewsArticleService.java (Dòng 89-122)",
        "Phương thức updateNews trong NewsArticleService: Cập nhật bài viết và ghi nhận thời gian modifiedDate.",
        "Bài viết cập nhật tiêu đề mới thành công, hiển thị chuẩn xác trên giao diện."
    ),
    (
        "noteTT/screenshots_a2/20_a2_news_delete_success.png",
        "Chức năng 4.7: Xóa bài viết Tin tức thành công",
        "Bấm nút 'Delete' tại bài viết, xác nhận hộp thoại cảnh báo. Backend gọi DELETE /api/news/{id} xóa bài viết và xóa các liên kết trong bảng liên kết NewsTag.",
        "Bài viết biến mất khỏi giao diện, danh sách cập nhật sạch sẽ.",
        "backend/.../service/NewsArticleService.java (Dòng 125-128) & frontend/src/pages/NewsList.jsx (Dòng 107-115)",
        "Phương thức handleDelete và deleteNews: Xóa bài viết khỏi cơ sở dữ liệu.",
        "Bài viết được xóa thành công khỏi database và giao diện."
    ),

    # Nhóm 5: Profile & Lịch sử tin tức (Staff)
    (
        "noteTT/screenshots_a2/21_a2_staff_profile.png",
        "Chức năng 5.1: Giao diện xem và cập nhật Hồ sơ cá nhân của Staff (Profile)",
        "Staff bấm menu 'My Profile' (/profile). Frontend gọi GET /api/auth/profile. Form hiển thị thông tin Name, Email và ô nhập mật khẩu mới. Bấm 'Save Profile' sẽ gọi PUT /api/auth/profile để cập nhật.",
        "Form hiển thị đúng tên và email của nhân viên đang đăng nhập, cho phép chỉnh sửa hồ sơ thuận tiện.",
        "backend/.../controller/AuthController.java (Dòng 22-29) & frontend/src/pages/Profile.jsx (Dòng 5-76)",
        "Phương thức profile và updateProfile trong AuthController: Quản lý hồ sơ cá nhân của nhân viên đang đăng nhập.",
        "Trang hồ sơ cá nhân hiển thị chuẩn xác thông tin tài khoản và cho phép chỉnh sửa."
    ),
    (
        "noteTT/screenshots_a2/22_a2_staff_my_history.png",
        "Chức năng 5.2: Giao diện xem Lịch sử các bài viết do chính Staff tạo (My History)",
        "Staff bấm menu 'My News History' (/my-history). Frontend gọi GET /api/news/created-by/{userId}. Backend lọc danh sách bài viết theo createdById từ session.",
        "Bảng hiển thị toàn bộ các bài viết do chính nhân viên này xuất bản, gồm Article ID, Title, Category, Status và Created Date chi tiết.",
        "backend/.../controller/NewsArticleController.java (Dòng 34-38) & frontend/src/pages/MyHistory.jsx (Dòng 6-88)",
        "Phương thức getCreatedBy trong NewsArticleController: Trả về danh sách bài báo do tác giả đăng nhập xuất bản.",
        "Lịch sử bài viết hiển thị đầy đủ danh sách tin tức do chính tài khoản Staff đang đăng nhập tạo ra."
    ),

    # Nhóm 6: Quản lý Tài khoản (Admin)
    (
        "noteTT/screenshots_a2/23_a2_admin_login_success.png",
        "Chức năng 6.1: Đăng nhập thành công với vai trò Quản trị viên (Admin - Role 1)",
        "Đăng xuất tài khoản Staff và đăng nhập bằng tài khoản admin@funews.com / @1. Backend xác thực và cấp quyền Admin.",
        "Navbar hiển thị thông tin 'Administrator', badge màu đỏ 'Admin', và xuất hiện menu quản trị độc quyền: Accounts Management.",
        "frontend/src/components/AppNavBar.jsx (Dòng 37-41 & 49-52)",
        "Điều kiện render user?.accountRole === 1: Mở khóa quyền hạn cao nhất cho Quản trị viên.",
        "Đăng nhập Admin thành công, thanh điều hướng xuất hiện menu Accounts Management."
    ),
    (
        "noteTT/screenshots_a2/24_a2_account_list.png",
        "Chức năng 6.2: Màn hình Quản lý Tài khoản người dùng hệ thống (Account Management)",
        "Admin truy cập /accounts. Frontend gọi GET /api/accounts. Backend Interceptor xác thực role=1, trả về danh sách tất cả tài khoản.",
        "Bảng hiển thị toàn bộ tài khoản gồm ID, Full Name, Email, Role badge (Admin màu đỏ, Staff màu xanh) và các nút thao tác Edit, Delete.",
        "backend/.../controller/SystemAccountController.java (Dòng 24-27) & frontend/src/pages/AccountList.jsx (Dòng 6-60)",
        "Phương thức getAll trong SystemAccountController: Lấy danh sách toàn bộ tài khoản người dùng cho Admin.",
        "Bảng tài khoản hiển thị đầy đủ thông tin nhân viên và quản trị viên trong hệ thống."
    ),
    (
        "noteTT/screenshots_a2/25_a2_account_create_modal.png",
        "Chức năng 6.3: Mở Popup Modal thêm mới Tài khoản người dùng",
        "Admin bấm '+ Create Account'. Modal mở ra với các trường: Full Name, Email Address, Password và dropdown chọn Role (Staff / Admin).",
        "Modal hiển thị tiêu đề 'Create Account', các trường nhập liệu sẵn sàng cho việc cấp phát tài khoản mới.",
        "frontend/src/pages/AccountList.jsx (Dòng 50-59 & 148-210)",
        "Hàm handleOpenCreate trong AccountList.jsx: Mở hộp thoại cấp tài khoản mới.",
        "Modal thêm mới tài khoản mở ra với các trường thông tin và chọn vai trò linh hoạt."
    ),
    (
        "noteTT/screenshots_a2/26_a2_account_create_success.png",
        "Chức năng 6.4: Thêm mới Tài khoản người dùng thành công",
        "Nhập Full Name 'Thai Assistant', Email 'assistant_thai@funews.com', mật khẩu 'secret123', Role 'Staff' rồi bấm 'Create Account'. Backend JPA lưu tài khoản mới vào SQL Server.",
        "Modal đóng lại, dòng tài khoản mới assistant_thai@funews.com xuất hiện ngay trong bảng với badge vai trò Staff.",
        "backend/.../service/SystemAccountService.java (Dòng 38-59)",
        "Phương thức createAccount trong SystemAccountService: Tạo tài khoản người dùng mới có mã hóa mật khẩu và tự sinh ID.",
        "Tài khoản mới được tạo thành công trong SQL Server và hiển thị trên bảng quản trị."
    ),
    (
        "noteTT/screenshots_a2/27_a2_account_validation_email_duplicate.png",
        "Chức năng 6.5: Bắt lỗi Validation chống trùng lặp Email tài khoản",
        "Thử bấm '+ Create Account' và nhập lại email assistant_thai@funews.com vừa tạo. Backend kiểm tra findByAccountEmail, phát hiện trùng lặp nên ném BusinessRuleViolationException('Email already exists').",
        "Xuất hiện thanh Alert màu đỏ cảnh báo: 'Email already exists: assistant_thai@funews.com'. Tài khoản trùng lặp bị từ chối tạo.",
        "backend/.../service/SystemAccountService.java (Dòng 40-42)",
        "Kiểm tra trùng lặp email duy nhất trong SystemAccountService: Đảm bảo định danh duy nhất cho mỗi tài khoản.",
        "Hệ thống chặn tạo tài khoản trùng email và hiển thị thông báo lỗi rõ ràng."
    ),
    (
        "noteTT/screenshots_a2/28_a2_account_update_modal.png",
        "Chức năng 6.6: Mở Popup Modal chỉnh sửa Tài khoản người dùng",
        "Bấm nút 'Edit' tại tài khoản assistant_thai@funews.com. Modal mở ra với thông tin điền sẵn. Admin nâng cấp Role từ Staff thành Administrator (Role 1).",
        "Modal hiển thị tiêu đề 'Edit Account', ô Password cho phép để trống nếu không muốn đổi mật khẩu.",
        "frontend/src/pages/AccountList.jsx (Dòng 62-72)",
        "Hàm handleOpenEdit: Nạp dữ liệu tài khoản vào form chỉnh sửa.",
        "Modal chỉnh sửa tài khoản cho phép Admin thăng cấp vai trò hoặc cập nhật thông tin."
    ),
    (
        "noteTT/screenshots_a2/29_a2_account_update_success.png",
        "Chức năng 6.7: Cập nhật Tài khoản người dùng thành công",
        "Bấm 'Save Changes', request PUT gửi lên backend cập nhật thông tin trong SQL Server.",
        "Tại dòng tài khoản assistant_thai@funews.com, badge vai trò đã đổi sang màu đỏ với nhãn 'Admin'.",
        "backend/.../service/SystemAccountService.java (Dòng 61-76)",
        "Phương thức updateAccount trong SystemAccountService: Cập nhật thông tin tài khoản người dùng.",
        "Tài khoản được nâng cấp lên vai trò Admin thành công, badge màu đỏ hiển thị tức thì."
    ),
    (
        "noteTT/screenshots_a2/30_a2_account_delete_prevent_in_use.png",
        "Chức năng 6.8: Chặn xóa Tài khoản đã từng tạo bài viết tin tức (Ràng buộc nghiệp vụ)",
        "Thử bấm 'Delete' tại tài khoản staff@funews.com (tài khoản này đã tạo bài viết tin tức trong hệ thống). Backend kiểm tra newsArticleRepository.existsByCreatedById(id). Vì đã tạo tin, backend ném BusinessRuleViolationException.",
        "Phía trên bảng xuất hiện Alert màu đỏ với thông báo rõ ràng: 'Cannot delete Account because it is the creator of existing news articles'. Tài khoản được bảo toàn.",
        "backend/.../service/SystemAccountService.java (Dòng 78-84)",
        "Phương thức deleteAccount trong SystemAccountService: Thực thi quy tắc bảo vệ tác giả bài báo.",
        "Hệ thống chặn xóa tài khoản đã từng tạo bài viết tin tức, bảo vệ lịch sử tác giả."
    ),
    (
        "noteTT/screenshots_a2/31_a2_account_delete_success.png",
        "Chức năng 6.9: Xóa Tài khoản người dùng thành công",
        "Bấm 'Delete' tại tài khoản assistant_thai@funews.com (chưa từng tạo bài viết nào), xác nhận hộp thoại cảnh báo. Backend xóa tài khoản khỏi SQL Server.",
        "Tài khoản biến mất khỏi bảng danh sách, cơ sở dữ liệu được làm sạch thành công.",
        "backend/.../service/SystemAccountService.java (Dòng 83) & frontend/src/pages/AccountList.jsx (Dòng 92-100)",
        "Phương thức handleDelete và deleteAccount: Xóa tài khoản hợp lệ khỏi database.",
        "Tài khoản chưa có bài viết được xóa thành công sau khi xác nhận hộp thoại."
    ),
    (
        "noteTT/screenshots_a2/32_a2_account_search.png",
        "Chức năng 6.10: Tìm kiếm Tài khoản người dùng theo tên hoặc email",
        "Nhập từ khóa 'staff' vào ô tìm kiếm. Backend lọc bằng query findByAccountNameContainingIgnoreCaseOrAccountEmailContainingIgnoreCase.",
        "Bảng chỉ hiển thị các tài khoản có tên hoặc email chứa từ khóa 'staff'.",
        "backend/.../repository/SystemAccountRepository.java (Dòng 12) & backend/.../service/SystemAccountService.java (Dòng 26-30)",
        "Phương thức getAllAccounts(search): Tìm kiếm tài khoản đa trường không phân biệt hoa thường.",
        "Tìm kiếm tài khoản theo tên hoặc email hoạt động chính xác qua API backend."
    ),
    (
        "noteTT/screenshots_a2/33_a2_logout_success.png",
        "Chức năng 6.11: Đăng xuất an toàn khỏi hệ thống (Logout)",
        "Admin bấm nút 'Logout'. Frontend gọi POST /api/auth/logout. Backend gọi session.invalidate() hủy toàn bộ session, frontend xóa user khỏi AuthContext và chuyển hướng về /news.",
        "Navbar trở về trạng thái Public, các menu quản trị biến mất hoàn toàn và xuất hiện nút 'Login (Staff / Admin)'.",
        "backend/.../controller/AuthController.java (Dòng 31-33) & frontend/src/components/AppNavBar.jsx (Dòng 9-13)",
        "Phương thức logout trong AuthController: Hủy bỏ phiên làm việc HttpSession an toàn.",
        "Đăng xuất thành công, phiên làm việc kết thúc và giao diện trở về trạng thái độc giả Public."
    ),
]

def build_combined_detailed_doc():
    doc = Document()
    for s in doc.sections:
        s.top_margin = Inches(0.8)
        s.bottom_margin = Inches(0.8)
        s.left_margin = Inches(0.8)
        s.right_margin = Inches(0.8)

    # 1. Title Banner
    add_header_banner(
        doc,
        "BÁO CÁO TOÀN DIỆN KẾT QUẢ TRIỂN KHAI VÀ KIỂM THỬ HỆ THỐNG",
        "HỌC PHẦN: SBA301 - XÂY DỰNG ỨNG DỤNG DOANH NGHIỆP | ĐẠI HỌC FPT"
    )

    # Student metadata table
    tbl_meta = doc.add_table(rows=6, cols=2)
    tbl_meta.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta_info = [
        ("Họ và tên sinh viên:", "Trần Quốc Thái"),
        ("Mã số sinh viên (MSSV):", "CE190243"),
        ("Lớp học phần:", "SE1910"),
        ("Dự án 1 (Assignment 01):", "TranQuocThai_SE1910 (ReactJS SPA & LocalStorage Mock Data)"),
        ("Dự án 2 (Assignment 02):", "A2TranQuocThai_SE1910 (Fullstack RESTful API, Spring Boot 3, JPA, SQL Server, ReactJS)"),
        ("Hệ thống triển khai:", "FUNewsManagementSystem - Quản lý Tin tức Đại học FPT"),
    ]
    for i, (k, v) in enumerate(meta_info):
        r = tbl_meta.rows[i]
        c0, c1 = r.cells[0], r.cells[1]
        c0.width, c1.width = Inches(2.3), Inches(4.2)
        set_cell_background(c0, "F1F5F9")
        set_cell_background(c1, "FFFFFF")
        set_cell_margins(c0, 50, 50, 100, 100)
        set_cell_margins(c1, 50, 50, 100, 100)
        p0 = c0.paragraphs[0]
        r0 = p0.add_run(k)
        r0.font.bold = True
        r0.font.size = Pt(9.5)
        p1 = c1.paragraphs[0]
        r1 = p1.add_run(v)
        r1.font.size = Pt(9.5)

    doc.add_paragraph().paragraph_format.space_after = Pt(14)

    # =========================================================================
    # PHẦN 1: ASSIGNMENT 01 (TranQuocThai_SE1910)
    # =========================================================================
    h_a1_banner = doc.add_heading("PHẦN 1: DỰ ÁN ASSIGNMENT 01 — TranQuocThai_SE1910", level=1)
    for r in h_a1_banner.runs:
        r.font.color.rgb = PRIMARY_COLOR
        r.font.size = Pt(15)
        r.font.bold = True

    add_callout(
        doc,
        "Đặc thù kỹ thuật của Assignment 01",
        "Ứng dụng Client-Side SPA thuần túy phát triển bằng ReactJS 19 và Vite 8. "
        "Toàn bộ dữ liệu được lưu trữ và đồng bộ bền vững trong LocalStorage của trình duyệt; "
        "chưa yêu cầu backend REST API hay database SQL Server theo chuẩn phân kỳ môn học SBA301."
    )

    p_a1_ov = doc.add_paragraph()
    p_a1_ov.paragraph_format.space_after = Pt(6)
    p_a1_ov.add_run("1.1 Mục tiêu đề tài và yêu cầu:\n").font.bold = True
    p_a1_ov.add_run(
        "Xây dựng giao diện quản trị tin tức FUNewsManagementSystem gồm 5 màn hình: Dashboard, Category, News, Users, Settings. "
        "Yêu cầu thực hiện đầy đủ CRUD và tìm kiếm cho Thể loại, Tin tức (liên kết Category) và Người dùng (Admin/Staff), "
        "form tạo/sửa dùng Modal popup, xóa qua Confirm dialog và chặn xóa Category khi đang có bài viết liên kết."
    )

    p_a1_sol = doc.add_paragraph()
    p_a1_sol.paragraph_format.space_after = Pt(8)
    p_a1_sol.add_run("1.2 Giải pháp kỹ thuật áp dụng:\n").font.bold = True
    p_a1_sol.add_run(
        "• Quản lý trạng thái & Persistence: Custom Hook useLocalStorage kết hợp useState và useEffect.\n"
        "• Component hóa: Header, Sidebar, Modal, ConfirmDialog, StatusBadge tái sử dụng cao.\n"
        "• Ràng buộc toàn vẹn dữ liệu: news.some(...) chặn xóa Category đang có bài viết.\n"
        "• Giao diện UI/UX: Hệ thống CSS Design Tokens, responsive chuẩn mực."
    )

    # Add 39 detailed steps of A1
    h_a1_steps = doc.add_heading("1.3 Kết quả kiểm thử chi tiết và mã nguồn Assignment 01 (39 Bước)", level=2)
    for r in h_a1_steps.runs:
        r.font.color.rgb = SECONDARY_COLOR
        r.font.size = Pt(13)
        r.font.bold = True

    # Import steps data from create_docs_a1
    # Using existing step list
    # (Image, Title, Works, Result, Loc, Fn)
    # We can iterate steps in create_docs_a1
    # Let's rebuild them quickly
    for item in create_docs_a1_steps_data:
        add_detailed_step(doc, item[0], item[1], item[2], item[3], item[4], item[5])

    # A1 Test Matrix
    doc.add_page_break()
    h_a1_matrix = doc.add_heading("1.4 Bảng tổng kết Ma trận kiểm thử Assignment 01", level=2)
    for r in h_a1_matrix.runs:
        r.font.color.rgb = SECONDARY_COLOR
        r.font.size = Pt(13)
        r.font.bold = True
    add_a1_matrix_table(doc)

    # =========================================================================
    # PHẦN 2: ASSIGNMENT 02 (A2TranQuocThai_SE1910)
    # =========================================================================
    doc.add_page_break()
    h_a2_banner = doc.add_heading("PHẦN 2: DỰ ÁN ASSIGNMENT 02 — A2TranQuocThai_SE1910", level=1)
    for r in h_a2_banner.runs:
        r.font.color.rgb = PRIMARY_COLOR
        r.font.size = Pt(15)
        r.font.bold = True

    add_callout(
        doc,
        "Đặc thù kiến trúc 3 lớp của Assignment 02",
        "Hệ thống Fullstack hoàn chỉnh theo kiến trúc 3 lớp (3-Tier Enterprise Architecture): "
        "Backend sử dụng Spring Boot 3 (Java 21), Hibernate JPA kết nối cơ sở dữ liệu thật SQL Server; "
        "Frontend sử dụng ReactJS, React-Bootstrap và Axios Client qua Vite Proxy /api. "
        "Xác thực và phân quyền dựa trên HttpSession và HandlerInterceptor.",
        bg_hex="FEF3C7",
        title_color=ACCENT_COLOR
    )

    p_a2_ov = doc.add_paragraph()
    p_a2_ov.paragraph_format.space_after = Pt(6)
    p_a2_ov.add_run("2.1 Đề tài Assignment 02 làm về gì và yêu cầu chi tiết:\n").font.bold = True
    p_a2_ov.add_run(
        "Assignment 02 nâng cấp hệ thống quản lý tin tức FUNews thành một dịch vụ Web Service RESTful kết nối cơ sở dữ liệu quan hệ SQL Server. "
        "Đề bài quy định rõ ràng 3 nhóm phân quyền người dùng:\n"
        "1. Độc giả công khai (Public): Được phép xem và tìm kiếm danh sách bài viết tin tức kích hoạt (Active).\n"
        "2. Nhân viên (Staff - Role 2): Được phép CRUD và tìm kiếm Thể loại (Category), CRUD và tìm kiếm Tin tức (News Articles kèm chọn nhiều Tags), "
        "xem/chỉnh sửa Hồ sơ cá nhân (My Profile), và xem Lịch sử các bài viết do chính mình xuất bản (My News History).\n"
        "3. Quản trị viên (Administrator - Role 1): Quản lý toàn bộ danh sách tài khoản người dùng hệ thống (System Accounts - CRUD & Search).\n"
        "4. Quy tắc nghiệp vụ bắt buộc (Business Rules): Chặn tuyệt đối việc xóa Category nếu đang có bài viết liên kết; "
        "chặn xóa Account nếu tài khoản đó là tác giả (creator) của bất kỳ bài viết tin tức nào; chặn trùng email tài khoản; "
        "tự động gán tác giả và thời gian tạo từ Session đăng nhập."
    )

    p_a2_sol = doc.add_paragraph()
    p_a2_sol.paragraph_format.space_after = Pt(8)
    p_a2_sol.add_run("2.2 Giải pháp kỹ thuật áp dụng trong A2TranQuocThai_SE1910:\n").font.bold = True
    p_a2_sol.add_run(
        "• Kiến trúc Backend 3 lớp (Controller → Service → Repository): Tách biệt hoàn toàn tầng tiếp nhận HTTP, xử lý nghiệp vụ và truy vấn dữ liệu.\n"
        "• Quản lý phiên và bảo mật (Session-based Auth & Interceptor): Sử dụng AccessConfig implements WebMvcConfigurer đăng ký HandlerInterceptor kiểm tra phiên đăng nhập và mã vai trò tự động cho mọi request /api/**.\n"
        "• Spring Data JPA & SQL Server: Sử dụng Hibernate ORM quản lý bảng Category, NewsArticle, SystemAccount, Tag và bảng quan hệ nhiều-nhiều NewsTag. Tự động sinh khóa chính và ngày tạo/sửa.\n"
        "• Xử lý ngoại lệ tập trung (Global Exception Handling): ApiExceptionHandler bắt BusinessRuleViolationException và ResourceNotFoundException, trả về mã lỗi HTTP 400/404 với thông điệp JSON chuẩn mực.\n"
        "• Frontend ReactJS & Axios: Sử dụng React Context (AuthContext) quản lý phiên, React Router DOM điều hướng SPA, React-Bootstrap xây dựng giao diện responsive và Axios Proxy /api tới port 8082."
    )

    # Add 33 detailed steps of A2
    h_a2_steps = doc.add_heading("2.3 Kết quả kiểm thử chi tiết và mã nguồn Assignment 02 (33 Bước)", level=2)
    for r in h_a2_steps.runs:
        r.font.color.rgb = SECONDARY_COLOR
        r.font.size = Pt(13)
        r.font.bold = True

    for item in a2_steps:
        add_detailed_step(doc, item[0], item[1], item[2], item[3], item[4], item[5])

    # A2 Test Matrix
    doc.add_page_break()
    h_a2_matrix = doc.add_heading("2.4 Bảng tổng kết Ma trận kiểm thử Assignment 02", level=2)
    for r in h_a2_matrix.runs:
        r.font.color.rgb = SECONDARY_COLOR
        r.font.size = Pt(13)
        r.font.bold = True
    add_a2_matrix_table(doc)

    out_file = os.path.abspath("noteTT/BAO_CAO_CHI_TIET_SBA301_TranQuocThai_SE1910.docx")
    doc.save(out_file)
    print(f"Combined detailed document created at: {out_file}")

    # Also try to update the original filename if unlocked
    try:
        doc.save(os.path.abspath("noteTT/BAO_CAO_CHI_TIET_ASSIGNMENT01_TranQuocThai_SE1910.docx"))
        print("Updated original file name as well!")
    except Exception as e:
        print(f"Original file is locked by Word: {e}")

def build_condensed_summary_doc():
    doc = Document()
    for s in doc.sections:
        s.top_margin = Inches(0.8)
        s.bottom_margin = Inches(0.8)
        s.left_margin = Inches(0.8)
        s.right_margin = Inches(0.8)

    # 1. Title Banner
    add_header_banner(
        doc,
        "TÀI LIỆU TÓM TẮT TRỌNG TÂM BÁO CÁO KỸ THUẬT SBA301",
        "TỔNG HỢP TOÀN BỘ MINH CHỨNG & KẾT QUẢ ĐÚNG (ASSIGNMENT 01 & ASSIGNMENT 02)"
    )

    # Student metadata
    p_meta = doc.add_paragraph()
    p_meta.paragraph_format.space_after = Pt(10)
    p_meta.add_run("Sinh viên: ").font.bold = True
    p_meta.add_run("Trần Quốc Thái (MSSV: CE190243 - Lớp: SE1910)\n")
    p_meta.add_run("Môn học: ").font.bold = True
    p_meta.add_run("SBA301 - Xây dựng ứng dụng Web Doanh nghiệp\n")
    p_meta.add_run("Mục đích tài liệu: ").font.bold = True
    p_meta.add_run(
        "Tài liệu tóm tắt nhanh toàn bộ nội dung trọng tâm của file Báo cáo chi tiết, "
        "giữ nguyên toàn bộ ảnh chụp minh chứng thực tế của tất cả các chức năng và rút gọn phần giải thích kiểm thử "
        "vào đúng 1 câu cốt lõi giúp người đọc nắm bắt toàn bộ luồng hệ thống một cách nhanh nhất."
    )

    # =========================================================================
    # PHẦN 1: TÓM TẮT ASSIGNMENT 01
    # =========================================================================
    h1 = doc.add_heading("PHẦN 1: TÓM TẮT ASSIGNMENT 01 — TranQuocThai_SE1910", level=1)
    for r in h1.runs:
        r.font.color.rgb = PRIMARY_COLOR
        r.font.size = Pt(14)
        r.font.bold = True

    add_callout(
        doc,
        "Tóm tắt kiến trúc Assignment 01",
        "Single Page Application xây dựng bằng ReactJS 19 + Vite 8. "
        "Dữ liệu mock được lưu trữ và đồng bộ bền vững trong trình duyệt qua LocalStorage (không dùng backend/DB thật). "
        "Bảo đảm đầy đủ CRUD Category, News, Users, modal popup, confirm dialog, chặn xóa Category đang có tin tức."
    )

    for item in create_docs_a1_summary_data:
        add_summary_step(doc, item[0], item[1], item[2])

    # =========================================================================
    # PHẦN 2: TÓM TẮT ASSIGNMENT 02
    # =========================================================================
    doc.add_page_break()
    h2 = doc.add_heading("PHẦN 2: TÓM TẮT ASSIGNMENT 02 — A2TranQuocThai_SE1910", level=1)
    for r in h2.runs:
        r.font.color.rgb = PRIMARY_COLOR
        r.font.size = Pt(14)
        r.font.bold = True

    add_callout(
        doc,
        "Tóm tắt kiến trúc Assignment 02",
        "Kiến trúc 3 lớp hoàn chỉnh: Backend Spring Boot 3 RESTful API kết nối cơ sở dữ liệu thật SQL Server qua JPA; "
        "Frontend ReactJS kết nối qua Axios. Phân quyền chặt chẽ bằng HttpSession & HandlerInterceptor: "
        "Public đọc tin Active; Staff quản lý Category, News (chọn Tags), Profile, History; Admin quản lý Accounts. "
        "Ràng buộc bắt buộc: Chặn xóa Category hoặc Account đang có bài báo liên kết.",
        bg_hex="FEF3C7",
        title_color=ACCENT_COLOR
    )

    for item in a2_steps:
        # item[0]: img, item[1]: title, item[6]: short summary
        add_summary_step(doc, item[0], item[1], item[6])

    out_file = os.path.abspath("noteTT/TOM_TAT_BAO_CAO_SBA301_TranQuocThai_SE1910.docx")
    doc.save(out_file)
    print(f"Condensed summary document created at: {out_file}")

def add_a1_matrix_table(doc):
    tbl = doc.add_table(rows=1, cols=4)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    headers = ["Mã Yêu Cầu", "Nội Dung Yêu Cầu Đề Bài", "Thành Phần Code & Bằng Chứng UI", "Kết Quả Kiểm Thử"]
    col_w = [Inches(1.1), Inches(2.2), Inches(2.2), Inches(1.1)]
    for j, h_text in enumerate(headers):
        tbl.rows[0].cells[j].width = col_w[j]
        set_cell_background(tbl.rows[0].cells[j], "1E3A8A")
        set_cell_margins(tbl.rows[0].cells[j], 80, 80, 100, 100)
        p = tbl.rows[0].cells[j].paragraphs[0]
        r = p.add_run(h_text)
        r.font.bold = True
        r.font.size = Pt(9)
        r.font.color.rgb = RGBColor(255, 255, 255)

    matrix_data = [
        ("R01", "Single Page Application với ReactJS và Vite", "package.json, vite.config.js, App.jsx", "ĐẠT (PASS 100%)"),
        ("R02", "Trang đăng nhập có validation rỗng và sai credential", "src/pages/LoginPage.jsx", "ĐẠT (PASS 100%)"),
        ("R03", "Đăng nhập tài khoản Admin/Admin vào Dashboard", "Auth state trong App.jsx, LoginPage.jsx", "ĐẠT (PASS 100%)"),
        ("R04", "Logo thương hiệu tạo bằng công cụ AI", "public/assets/funews-logo.png, Header, Sidebar", "ĐẠT (PASS 100%)"),
        ("R05", "Bố cục Header và Sidebar dùng chung", "src/components/Header.jsx, Sidebar.jsx", "ĐẠT (PASS 100%)"),
        ("R06", "Menu thanh điều hướng đủ 5 mục chức năng", "Sidebar.jsx (Dashboard, Category, News, Users, Settings)", "ĐẠT (PASS 100%)"),
        ("R07", "CRUD & Search cho Thể loại (Category)", "src/pages/CategoryPage.jsx, StatusBadge.jsx", "ĐẠT (PASS 100%)"),
        ("R08", "CRUD & Search cho Tin tức (News)", "src/pages/NewsPage.jsx (liên kết Category)", "ĐẠT (PASS 100%)"),
        ("R09", "CRUD & Search cho Tài khoản Người dùng (Users)", "src/pages/UsersPage.jsx (Role Admin/Staff)", "ĐẠT (PASS 100%)"),
        ("R10", "Giao diện Thêm mới & Cập nhật dùng Popup Modal", "src/components/Modal.jsx", "ĐẠT (PASS 100%)"),
        ("R11", "Giao diện Xóa có hộp thoại xác nhận (Confirm Dialog)", "src/components/ConfirmDialog.jsx", "ĐẠT (PASS 100%)"),
        ("R12", "Ràng buộc toàn vẹn dữ liệu (Chặn xóa Category có News)", "handleDelete trong CategoryPage.jsx", "ĐẠT (PASS 100%)"),
        ("R13", "Lưu trữ dữ liệu bền vững trong trình duyệt", "src/hooks/useLocalStorage.js", "ĐẠT (PASS 100%)"),
    ]
    for m_row in matrix_data:
        row = tbl.add_row()
        for c_idx, val in enumerate(m_row):
            cell = row.cells[c_idx]
            cell.width = col_w[c_idx]
            set_cell_background(cell, "F8FAFC" if c_idx == 0 else "FFFFFF")
            set_cell_margins(cell, 60, 60, 100, 100)
            p = cell.paragraphs[0]
            r = p.add_run(val)
            r.font.size = Pt(8.5)
            if c_idx == 3:
                r.font.bold = True
                r.font.color.rgb = SUCCESS_COLOR
            elif c_idx == 0:
                r.font.bold = True
                r.font.color.rgb = PRIMARY_COLOR

def add_a2_matrix_table(doc):
    tbl = doc.add_table(rows=1, cols=4)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    headers = ["Mã Yêu Cầu", "Nội Dung Yêu Cầu Đề Bài", "Thành Phần Code & Bằng Chứng UI", "Kết Quả Kiểm Thử"]
    col_w = [Inches(1.1), Inches(2.2), Inches(2.2), Inches(1.1)]
    for j, h_text in enumerate(headers):
        tbl.rows[0].cells[j].width = col_w[j]
        set_cell_background(tbl.rows[0].cells[j], "1E3A8A")
        set_cell_margins(tbl.rows[0].cells[j], 80, 80, 100, 100)
        p = tbl.rows[0].cells[j].paragraphs[0]
        r = p.add_run(h_text)
        r.font.bold = True
        r.font.size = Pt(9)
        r.font.color.rgb = RGBColor(255, 255, 255)

    matrix_data = [
        ("A2-01", "Kiến trúc 3 lớp Spring Boot RESTful Web Service", "Controller -> Service -> Repository, pom.xml", "ĐẠT (PASS 100%)"),
        ("A2-02", "Cơ sở dữ liệu quan hệ SQL Server qua Spring Data JPA", "Entities: Category, NewsArticle, SystemAccount, Tag", "ĐẠT (PASS 100%)"),
        ("A2-03", "Độc giả công khai xem và tìm kiếm tin tức Active", "/api/news/active, NewsList.jsx", "ĐẠT (PASS 100%)"),
        ("A2-04", "Xác thực phiên làm việc bằng HttpSession & Interceptor", "AccessConfig.java, AuthController.java, AuthContext.jsx", "ĐẠT (PASS 100%)"),
        ("A2-05", "Nhân viên (Staff) CRUD và tìm kiếm Thể loại", "CategoryController.java, CategoryService.java, CategoryList.jsx", "ĐẠT (PASS 100%)"),
        ("A2-06", "Ràng buộc nghiệp vụ: Chặn xóa Thể loại đang có bài báo", "CategoryService.deleteCategory -> existsByCategory", "ĐẠT (PASS 100%)"),
        ("A2-07", "Nhân viên (Staff) CRUD tin tức kèm quan hệ nhiều-nhiều Tags", "NewsArticleService.java, NewsList.jsx, NewsTag table", "ĐẠT (PASS 100%)"),
        ("A2-08", "Tự động gán tác giả từ Session và ngày tạo/ngày sửa", "NewsArticleController.create -> setCreatedById(session)", "ĐẠT (PASS 100%)"),
        ("A2-09", "Nhân viên xem và chỉnh sửa Hồ sơ cá nhân (My Profile)", "AuthController.updateProfile, Profile.jsx", "ĐẠT (PASS 100%)"),
        ("A2-10", "Nhân viên xem Lịch sử tin tức do chính mình xuất bản", "/api/news/created-by/{id}, MyHistory.jsx", "ĐẠT (PASS 100%)"),
        ("A2-11", "Quản trị viên (Admin) CRUD và tìm kiếm Tài khoản hệ thống", "SystemAccountController.java, SystemAccountService.java", "ĐẠT (PASS 100%)"),
        ("A2-12", "Ràng buộc nghiệp vụ: Chặn xóa Tài khoản đã tạo bài báo", "SystemAccountService.deleteAccount -> existsByCreatedById", "ĐẠT (PASS 100%)"),
        ("A2-13", "Ràng buộc nghiệp vụ: Chống trùng lặp Email tài khoản", "SystemAccountService.createAccount -> findByAccountEmail", "ĐẠT (PASS 100%)"),
        ("A2-14", "Form tạo/sửa dùng Modal Popup, xóa có Confirm xác nhận", "CategoryForm.jsx, NewsList.jsx, AccountList.jsx", "ĐẠT (PASS 100%)"),
    ]
    for m_row in matrix_data:
        row = tbl.add_row()
        for c_idx, val in enumerate(m_row):
            cell = row.cells[c_idx]
            cell.width = col_w[c_idx]
            set_cell_background(cell, "F8FAFC" if c_idx == 0 else "FFFFFF")
            set_cell_margins(cell, 60, 60, 100, 100)
            p = cell.paragraphs[0]
            r = p.add_run(val)
            r.font.size = Pt(8.5)
            if c_idx == 3:
                r.font.bold = True
                r.font.color.rgb = SUCCESS_COLOR
            elif c_idx == 0:
                r.font.bold = True
                r.font.color.rgb = PRIMARY_COLOR

# Load data for A1
create_docs_a1_steps_data = a1_steps
create_docs_a1_summary_data = []
for s in a1_steps:
    # Generate 1-sentence summary for A1: (img, title, result)
    create_docs_a1_summary_data.append((s[0], s[1], s[3]))

if __name__ == "__main__":
    build_combined_detailed_doc()
    build_condensed_summary_doc()
