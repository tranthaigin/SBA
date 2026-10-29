import os
import sys
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

if hasattr(sys.stdout, 'reconfigure'):
    sys.stdout.reconfigure(encoding='utf-8')

PRIMARY_COLOR = RGBColor(30, 58, 138)    # #1E3A8A - Deep Blue
SECONDARY_COLOR = RGBColor(2, 132, 199)  # #0284C7 - Blue
ACCENT_COLOR = RGBColor(180, 83, 9)      # #B45309 - Amber
DARK_TEXT = RGBColor(31, 41, 55)         # #1F2937 - Dark Gray
MUTED_TEXT = RGBColor(107, 114, 128)     # #6B7280 - Light Gray
SUCCESS_COLOR = RGBColor(22, 101, 52)    # Green

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=140, right=140):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def add_header_banner(doc, title_text, subtitle_text, bg_hex="1E3A8A"):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    set_cell_background(cell, bg_hex)
    set_cell_margins(cell, top=200, bottom=200, left=250, right=250)
    
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run1 = p.add_run(title_text + "\n")
    run1.font.bold = True
    run1.font.size = Pt(16)
    run1.font.color.rgb = RGBColor(255, 255, 255)
    
    run2 = p.add_run(subtitle_text)
    run2.font.size = Pt(11)
    run2.font.color.rgb = RGBColor(224, 242, 254)
    
    p_after = doc.add_paragraph()
    p_after.paragraph_format.space_after = Pt(12)

def add_callout(doc, title, text, bg_hex="EFF6FF", title_color=PRIMARY_COLOR):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    set_cell_background(cell, bg_hex)
    set_cell_margins(cell, top=120, bottom=120, left=180, right=180)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(4)
    r_title = p.add_run(f"📌 {title}\n")
    r_title.font.bold = True
    r_title.font.size = Pt(10.5)
    r_title.font.color.rgb = title_color
    
    r_text = p.add_run(text)
    r_text.font.size = Pt(10)
    r_text.font.color.rgb = DARK_TEXT
    
    p_sp = doc.add_paragraph()
    p_sp.paragraph_format.space_after = Pt(8)

def add_detailed_step(doc, image_path, step_title, how_it_works, expected_result, code_location, function_explanation):
    h = doc.add_heading(step_title, level=3)
    h.paragraph_format.space_before = Pt(14)
    h.paragraph_format.space_after = Pt(6)
    for r in h.runs:
        r.font.color.rgb = PRIMARY_COLOR
        r.font.size = Pt(11.5)
        r.font.bold = True

    if os.path.exists(image_path):
        p_img = doc.add_paragraph()
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img.paragraph_format.space_before = Pt(4)
        p_img.paragraph_format.space_after = Pt(4)
        p_img.add_run().add_picture(image_path, width=Inches(5.9))
        
        p_cap = doc.add_paragraph()
        p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_cap.paragraph_format.space_after = Pt(6)
        r_cap = p_cap.add_run(f"Hình: {step_title}")
        r_cap.font.italic = True
        r_cap.font.size = Pt(9.5)
        r_cap.font.color.rgb = MUTED_TEXT

    tbl = doc.add_table(rows=4, cols=2)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    
    col_widths = [Inches(1.8), Inches(4.7)]
    rows_data = [
        ("Cơ chế hoạt động:", how_it_works),
        ("Kết quả nhìn thấy là đúng:", expected_result),
        ("Vị trí code tương ứng:", code_location),
        ("Phương thức / Hàm xử lý:", function_explanation),
    ]
    
    for i, (label, content) in enumerate(rows_data):
        row = tbl.rows[i]
        c0, c1 = row.cells[0], row.cells[1]
        c0.width, c1.width = col_widths[0], col_widths[1]
        set_cell_background(c0, "F8FAFC")
        set_cell_margins(c0, 70, 70, 110, 110)
        p0 = c0.paragraphs[0]
        p0.paragraph_format.space_after = Pt(0)
        r0 = p0.add_run(label)
        r0.font.bold = True
        r0.font.size = Pt(9)
        r0.font.color.rgb = PRIMARY_COLOR
        
        set_cell_background(c1, "FFFFFF" if i % 2 == 0 else "F9FAFB")
        set_cell_margins(c1, 70, 70, 110, 110)
        p1 = c1.paragraphs[0]
        p1.paragraph_format.space_after = Pt(0)
        r1 = p1.add_run(content)
        r1.font.size = Pt(9)
        r1.font.color.rgb = DARK_TEXT

    doc.add_paragraph().paragraph_format.space_after = Pt(10)

def add_summary_step(doc, image_path, step_title, summary_desc):
    h = doc.add_heading(step_title, level=3)
    h.paragraph_format.space_before = Pt(12)
    h.paragraph_format.space_after = Pt(4)
    for r in h.runs:
        r.font.color.rgb = PRIMARY_COLOR
        r.font.size = Pt(11)
        r.font.bold = True

    if os.path.exists(image_path):
        p_img = doc.add_paragraph()
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img.paragraph_format.space_before = Pt(2)
        p_img.paragraph_format.space_after = Pt(4)
        p_img.add_run().add_picture(image_path, width=Inches(5.9))

    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    set_cell_background(cell, "F0FDF4")
    set_cell_margins(cell, 70, 70, 120, 120)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    r1 = p.add_run("✔️ Trọng tâm kết quả đúng: ")
    r1.font.bold = True
    r1.font.size = Pt(9.5)
    r1.font.color.rgb = SUCCESS_COLOR
    r2 = p.add_run(summary_desc)
    r2.font.size = Pt(9.5)
    r2.font.color.rgb = SUCCESS_COLOR

    doc.add_paragraph().paragraph_format.space_after = Pt(8)
